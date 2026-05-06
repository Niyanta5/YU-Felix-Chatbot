#!/usr/bin/env python3
"""
generate_docs.py

Utility script to convert markdown KT documents into .docx format
using python-docx.
"""
import os
import re
from docx import Document

def md_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    buffer = []
    for line in lines:
        stripped = line.rstrip('\n')
        # Heading detection
        m = re.match(r'^(#+)\s+(.*)', stripped)
        if m:
            # flush buffer before heading
            if buffer:
                doc.add_paragraph('\n'.join(buffer))
                buffer = []
            level = min(len(m.group(1)), 4)
            text = m.group(2)
            doc.add_heading(text, level=level)
            continue
        # Bullet list detection
        if re.match(r'^[\-\*+]\s+', stripped):
            if buffer:
                doc.add_paragraph('\n'.join(buffer))
                buffer = []
            text = re.sub(r'^[\-\*+]\s+', '', stripped)
            doc.add_paragraph(text, style='List Bullet')
            continue
        # Paragraph separation on blank line
        if stripped.strip() == '':
            if buffer:
                doc.add_paragraph('\n'.join(buffer))
                buffer = []
            continue
        # Regular paragraph text
        buffer.append(stripped)
    # flush remaining buffer
    if buffer:
        doc.add_paragraph('\n'.join(buffer))
    # Save the document
    doc.save(docx_path)

if __name__ == '__main__':
    base_dir = os.path.dirname(__file__)
    docs = ['Technical_Document.md', 'NonTechnical_Document.md']
    for md in docs:
        md_path = os.path.join(base_dir, md)
        docx_path = os.path.join(base_dir, md.replace('.md', '.docx'))
        if os.path.exists(md_path):
            print(f"Converting {md} -> {os.path.basename(docx_path)}")
            md_to_docx(md_path, docx_path)
        else:
            print(f"Markdown file not found: {md_path}")